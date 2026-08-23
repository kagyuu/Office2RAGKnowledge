"""テスト用の Office ファイルをコードで生成するヘルパ (P006 §3.1)。

バイナリをリポジトリに大量投入せずに済むよう、可能なものはすべてコードで作る。
**openpyxl で書き出したファイルは自動的にキャッシュ値を持たない**ため、
要求②(数式キャッシュ値)のテストデータがそのまま作れる。
"""

from __future__ import annotations

import io
import shutil
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from lxml import etree
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

M_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

__all__ = [
    "add_sheet_drawing",
    "inject_cached_values",
    "make_docx",
    "make_pptx",
    "make_xlsx",
    "zip_from_mapping",
]


def make_xlsx(path: Path, sheets: Mapping[str, Sequence[Sequence[Any]]], **kwargs: Any) -> Path:
    """シート名 → 2次元の値で ``.xlsx`` を作る。

    Keyword Args:
        merges: シート名 → 結合範囲文字列の並び (例: ``{"S1": ["A1:C1"]}``)
        hidden_sheets: 非表示にするシート名の集合
        hidden_rows: シート名 → 非表示にする行番号の並び
        number_formats: シート名 → {セル参照: 表示形式}
    """
    merges: Mapping[str, Iterable[str]] = kwargs.get("merges", {})
    hidden_sheets: set[str] = set(kwargs.get("hidden_sheets", ()))
    hidden_rows: Mapping[str, Iterable[int]] = kwargs.get("hidden_rows", {})
    number_formats: Mapping[str, Mapping[str, str]] = kwargs.get("number_formats", {})

    wb = Workbook()
    default = wb.active
    assert default is not None
    first = True
    for name, rows in sheets.items():
        ws = default if first else wb.create_sheet()
        ws.title = name
        first = False
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if value is None or value == "":
                    continue
                ws.cell(row=r, column=c, value=value)
        for ref in merges.get(name, ()):
            ws.merge_cells(ref)
        for row_idx in hidden_rows.get(name, ()):
            ws.row_dimensions[row_idx].hidden = True
        for ref, fmt in number_formats.get(name, {}).items():
            ws[ref].number_format = fmt
        if name in hidden_sheets:
            ws.sheet_state = "hidden"
    wb.save(path)
    return path


def inject_cached_values(path: Path, value: str = "999") -> Path:
    """数式セルに計算結果 (キャッシュ値) を注入する。

    openpyxl は数式セルに **空の ``<v>``** を書き出すため、その ``<v>`` の
    テキストを埋めることで「Excel が保存したファイル」を模す。
    新しい ``<v>`` を追加すると要素が2つになるため、既存のものを再利用する。
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        sheet_names = [n for n in names if n.startswith("xl/worksheets/sheet")]
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names:
                data = zin.read(name)
                if name in sheet_names:
                    root = etree.fromstring(data)
                    for cell in root.iter(f"{{{M_NS}}}c"):
                        if cell.find(f"{{{M_NS}}}f") is None:
                            continue
                        v = cell.find(f"{{{M_NS}}}v")
                        if v is None:
                            v = etree.SubElement(cell, f"{{{M_NS}}}v")
                        v.text = value
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
                zout.writestr(name, data)
    path.write_bytes(buf.getvalue())
    return path


def add_sheet_drawing(path: Path, sheet_index: int, drawing_xml: str) -> Path:
    """シートに drawing パートを注入する。

    openpyxl は図形を作れないため、ZIP に直接パートと関係を追加する。
    ``sheet_index`` は 0 起点 (``xl/worksheets/sheet1.xml`` が 0)。
    """
    drawing_no = sheet_index + 1
    drawing_part = f"xl/drawings/drawing{drawing_no}.xml"
    sheet_part = f"xl/worksheets/sheet{drawing_no}.xml"
    rels_part = f"xl/worksheets/_rels/sheet{drawing_no}.xml.rels"

    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names:
                data = zin.read(name)
                if name == "[Content_Types].xml":
                    root = etree.fromstring(data)
                    override = etree.SubElement(root, f"{{{CT_NS}}}Override")
                    override.set("PartName", f"/{drawing_part}")
                    override.set(
                        "ContentType",
                        "application/vnd.openxmlformats-officedocument.drawing+xml",
                    )
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
                elif name == sheet_part:
                    root = etree.fromstring(data)
                    if root.find(f"{{{M_NS}}}drawing") is None:
                        drawing = etree.SubElement(root, f"{{{M_NS}}}drawing")
                        drawing.set(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
                            "rIdDrawing1",
                        )
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
                elif name == rels_part:
                    continue  # 下で作り直す
                zout.writestr(name, data)

            zout.writestr(drawing_part, drawing_xml)
            rels = etree.Element(f"{{{REL_NS}}}Relationships")
            rel = etree.SubElement(rels, f"{{{REL_NS}}}Relationship")
            rel.set("Id", "rIdDrawing1")
            rel.set(
                "Type",
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing",
            )
            rel.set("Target", f"../drawings/drawing{drawing_no}.xml")
            zout.writestr(rels_part, etree.tostring(rels, xml_declaration=True, encoding="UTF-8"))
    path.write_bytes(buf.getvalue())
    return path


def make_docx(path: Path, blocks: Sequence[tuple[str, str]]) -> Path:
    """``(種別, テキスト)`` の並びから ``.docx`` を作る。

    種別: ``"h1"``〜``"h5"`` / ``"p"`` / ``"bullet"`` / ``"number"``
    """
    doc = DocxDocument()
    for kind, text in blocks:
        if kind.startswith("h") and kind[1:].isdigit():
            doc.add_heading(text, level=int(kind[1:]))
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(text, style="List Number")
        else:
            doc.add_paragraph(text)
    doc.save(path)
    return path


def make_pptx(path: Path, slides: Sequence[tuple[str | None, Sequence[str], str | None]]) -> Path:
    """``(タイトル, 本文の並び, ノート)`` の並びから ``.pptx`` を作る。"""
    prs = Presentation()
    for title, bullets, notes in slides:
        layout = prs.slide_layouts[5] if title is not None else prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        if title is not None and slide.shapes.title is not None:
            slide.shapes.title.text = title
        if bullets:
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
            frame = box.text_frame
            for i, line in enumerate(bullets):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                para.text = line
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def zip_from_mapping(path: Path, entries: Mapping[str, bytes | str]) -> Path:
    """辞書から ZIP を作る (SafeZipReader のテスト用)。"""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return path


def copy_to(src: Path, dst: Path) -> Path:
    """ファイルを複製する (元を壊さずに細工するため)。"""
    shutil.copy(src, dst)
    return dst


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def inject_docx_body_xml(path: Path, xml: str, *, index: int = -1) -> Path:
    """``word/document.xml`` の body に XML 断片を挿入する。

    python-docx はテキストボックス・変更履歴を作れないため、直接注入する。
    ``index=-1`` は sectPr の直前 (末尾) を意味する。
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in names:
                data = zin.read(name)
                if name == "word/document.xml":
                    root = etree.fromstring(data)
                    body = root.find(f"{{{W_NS}}}body")
                    assert body is not None
                    node = etree.fromstring(xml)
                    if index < 0:
                        body.insert(max(0, len(body) - 1), node)
                    else:
                        body.insert(index, node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
                zout.writestr(name, data)
    path.write_bytes(buf.getvalue())
    return path


def docx_textbox_xml(paragraphs: Sequence[str], *, name: str = "テキスト ボックス 1") -> str:
    """テキストボックスを含む段落の XML を返す (mc:Choice / mc:Fallback 付き)。"""
    inner = "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>" for t in paragraphs)
    return f"""<w:p xmlns:w="{W_NS}"
      xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
      xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
      xmlns:v="urn:schemas-microsoft-com:vml">
      <w:r>
        <mc:AlternateContent>
          <mc:Choice Requires="wps">
            <w:drawing><wp:inline><a:graphic><a:graphicData>
              <wps:wsp>
                <wps:cNvPr name="{name}"/>
                <wps:cNvSpPr txBox="1"/>
                <wps:txbx><w:txbxContent>{inner}</w:txbxContent></wps:txbx>
              </wps:wsp>
            </a:graphicData></a:graphic></wp:inline></w:drawing>
          </mc:Choice>
          <mc:Fallback>
            <w:pict><v:shape><v:textbox><w:txbxContent>{inner}</w:txbxContent>
            </v:textbox></v:shape></w:pict>
          </mc:Fallback>
        </mc:AlternateContent>
      </w:r>
    </w:p>"""


def docx_revision_xml(inserted: str, deleted: str) -> str:
    """変更履歴 (挿入・削除) を含む段落の XML を返す。"""
    return f"""<w:p xmlns:w="{W_NS}">
      <w:ins w:id="1" w:author="a"><w:r><w:t>{inserted}</w:t></w:r></w:ins>
      <w:del w:id="2" w:author="a"><w:r><w:delText>{deleted}</w:delText></w:r></w:del>
    </w:p>"""
